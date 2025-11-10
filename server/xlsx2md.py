import argparse, pathlib, re, sys
from openpyxl import load_workbook

def clean_cell(v):
    if v is None:
        return ""
    s = str(v)
    return re.sub(r"\s+", " ", s).strip()

def convert_one(xlsx_path: pathlib.Path, out_dir: pathlib.Path):
    try:
        wb = load_workbook(filename=str(xlsx_path), data_only=True, read_only=True)
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / (xlsx_path.stem + ".md")
        with out_path.open("w", encoding="utf-8") as f:
            f.write(f"# {xlsx_path.stem}\n\n")
            for ws in wb.worksheets:
                f.write(f"## {ws.title}\n\n")
                for row in ws.iter_rows(values_only=True):
                    cells = [clean_cell(c) for c in row]
                    # Línea tipo tabla simple: valor1 | valor2 | ...
                    f.write(" | ".join(cells) + "\n")
                f.write("\n")
        print(f"[ok] {xlsx_path.name} -> {out_path}")
    except Exception as e:
        print(f"[skip] {xlsx_path.name}: {e}", file=sys.stderr)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--dst", required=True)
    args = ap.parse_args()
    src = pathlib.Path(args.src)
    dst = pathlib.Path(args.dst)
    files = sorted(src.rglob("*.xlsx"))
    if not files:
        print("[warn] No se encontraron .xlsx")
    for f in files:
        convert_one(f, dst)

if __name__ == "__main__":
    main()
